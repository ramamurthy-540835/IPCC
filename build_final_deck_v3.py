#!/usr/bin/env python3
"""Rebuild IPCC AI GCP deck as an executive, visual-first presentation."""
import math
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette
NAVY   = RGBColor(0x0B, 0x2E, 0x4F)   # deep ocean
NAVY2  = RGBColor(0x11, 0x3F, 0x66)
TEAL   = RGBColor(0x0C, 0x7B, 0x93)   # sea teal
GREEN  = RGBColor(0x2E, 0x9E, 0x6B)   # leaf
AMBER  = RGBColor(0xE8, 0xA3, 0x3D)
RED    = RGBColor(0xD4, 0x5D, 0x5D)
BLUE   = RGBColor(0x5B, 0x84, 0xC4)
PURPLE = RGBColor(0x8B, 0x6F, 0xB8)
INK    = RGBColor(0x1C, 0x2B, 0x36)   # near-black text
GRAY   = RGBColor(0x5B, 0x6B, 0x79)   # body text
FAINT  = RGBColor(0x9A, 0xA7, 0xB2)
BG     = RGBColor(0xF8, 0xFA, 0xFB)   # page background
PANEL  = RGBColor(0xEE, 0xF1, 0xF4)   # neutral panel
TEALBG = RGBColor(0xE3, 0xF2, 0xEE)   # teal tint panel
NAVYBG = RGBColor(0xE8, 0xF0, 0xF7)
AMBERBG= RGBColor(0xFB, 0xF3, 0xE2)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
OFFW   = RGBColor(0xD9, 0xE4, 0xEC)   # light text on navy

F = "Segoe UI"
FS = "Segoe UI Semibold"
FL = "Segoe UI Light"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- helpers
def slide_new():
    return prs.slides.add_slide(BLANK)

def _set_alpha(shape, alpha):
    """alpha 0..1 on a solid fill"""
    sf = shape.fill._xPr.find(qn('a:solidFill'))
    clr = sf.find(qn('a:srgbClr'))
    a = clr.makeelement(qn('a:alpha'), {'val': str(int(alpha * 100000))})
    clr.append(a)

def shp(slide, kind, x, y, w, h, fill=None, line=None, lw=0.75, alpha=None, shadow_off=True):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
        if alpha is not None:
            _set_alpha(s, alpha)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    if shadow_off:
        s.shadow.inherit = False
    return s

def rrect(slide, x, y, w, h, fill=None, line=None, lw=0.75, radius=0.12, alpha=None):
    s = shp(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, lw, alpha)
    try:
        s.adjustments[0] = radius
    except Exception:
        pass
    return s

def tb(slide, x, y, w, h, text, size=12, color=GRAY, bold=False, align=PP_ALIGN.LEFT,
       font=F, anchor=MSO_ANCHOR.TOP, spc=None, line_sp=None, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_sp:
            p.line_spacing = line_sp
        r = p.add_run(); r.text = ln
        r.font.name = font; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = color
        if spc:
            r.font._rPr.set('spc', str(spc))
    return box

def multi(slide, x, y, w, h, runs_per_par, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_sp=None, space_after=None):
    """runs_per_par: list of paragraphs; each = list of (text,font,size,color,bold[,spc])"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, runs in enumerate(runs_per_par):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_sp: p.line_spacing = line_sp
        if space_after is not None: p.space_after = Pt(space_after)
        for run in runs:
            t, fnt, sz, col, bld = run[:5]
            r = p.add_run(); r.text = t
            r.font.name = fnt; r.font.size = Pt(sz); r.font.bold = bld
            r.font.color.rgb = col
            if len(run) > 5 and run[5]:
                r.font._rPr.set('spc', str(run[5]))
    return box

def connector(slide, x1, y1, x2, y2, color=FAINT, w=1.0, dash=None, arrow=False, both=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    c.shadow.inherit = False
    ln = c.line._get_or_add_ln()
    if dash:
        d = ln.makeelement(qn('a:prstDash'), {'val': dash}); ln.append(d)
    if arrow:
        t = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}); ln.append(t)
    if both:
        h = ln.makeelement(qn('a:headEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}); ln.append(h)
    return c

def page_bg(slide, color=BG):
    shp(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, fill=color)

def header(slide, kicker, headline, sub, page, accent=TEAL, hsize=27):
    page_bg(slide)
    shp(slide, MSO_SHAPE.RECTANGLE, 0.6, 0.62, 0.34, 0.055, fill=accent)
    tb(slide, 1.05, 0.52, 9.5, 0.3, kicker.upper(), 10.5, accent, True, spc=160)
    tb(slide, 0.6, 0.88, 12.1, 0.65, headline, hsize, INK, True, font=FS)
    tb(slide, 0.6, 1.5, 11.0, 0.35, sub, 12.5, GRAY)
    tb(slide, 12.55, 7.02, 0.55, 0.3, str(page), 10, FAINT, align=PP_ALIGN.RIGHT)

def footer(slide, text, accent=TEAL):
    connector(slide, 0.6, 6.98, 12.73, 6.98, color=RGBColor(0xDD, 0xE4, 0xE9), w=0.75)
    tb(slide, 0.6, 7.06, 11.6, 0.3, text, 9.5, FAINT)

def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def dot(slide, x, y, d, color, line=None, lw=1.0, alpha=None):
    return shp(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=color, line=line, lw=lw, alpha=alpha)

# ================================================================ SLIDE 1 — HERO
s = slide_new()
bgr = shp(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, fill=NAVY)
try:
    bgr.fill.gradient()
    stops = bgr.fill.gradient_stops
    stops[0].color.rgb = RGBColor(0x08, 0x22, 0x3C)
    stops[1].color.rgb = RGBColor(0x14, 0x45, 0x63)
    bgr.fill.gradient_angle = 35.0
except Exception:
    pass

# data constellation (right side): concentric orbit rings + nodes
cx, cy = 10.35, 3.9
for r in (0.9, 1.55, 2.2, 2.85):
    ring = shp(s, MSO_SHAPE.OVAL, cx - r, cy - r, 2*r, 2*r, fill=None, line=WHITE, lw=0.9)
    ln = ring.line._get_or_add_ln()
    sf = ln.find(qn('a:solidFill')); clr = sf.find(qn('a:srgbClr'))
    a = clr.makeelement(qn('a:alpha'), {'val': '14000'}); clr.append(a)
node_data = [(0.9, 200, TEAL, .16), (1.55, 340, GREEN, .18), (1.55, 120, AMBER, .15),
             (2.2, 30, TEAL, .20), (2.2, 255, GREEN, .14), (2.85, 155, AMBER, .17),
             (2.85, 305, TEAL, .13), (0.9, 60, WHITE, .12)]
for r, ang, col, d in node_data:
    nx = cx + r * math.cos(math.radians(ang)) - d/2
    ny = cy - r * math.sin(math.radians(ang)) - d/2
    dot(s, nx, ny, d, col)
core = dot(s, cx - 0.3, cy - 0.3, 0.6, GREEN)
dot(s, cx - 0.14, cy - 0.14, 0.28, WHITE, alpha=0.9)

tb(s, 0.75, 1.35, 8.0, 0.3, "IPCC CLIMATE AI ASSISTANT  ·  CEO / CIO BRIEFING", 12, RGBColor(0x6F, 0xD0, 0xC2), True, spc=200)
tb(s, 0.72, 1.8, 8.3, 2.0, "From thousands of pages\nto one grounded answer", 40, WHITE, True, font=FS, line_sp=1.02)
tb(s, 0.75, 3.55, 7.2, 0.8,
   "A conversational assistant that retrieves IPCC evidence first, explains it in plain language, and always shows its sources.",
   14, OFFW, line_sp=1.25)

# stat chips
chips = [("8", "IPCC reports"), ("10,184", "indexed passages"), ("5", "workflows"), ("LIVE", "on Cloud Run")]
chx = 0.75
for val, lab in chips:
    w = 1.95
    c = rrect(s, chx, 4.85, w, 1.0, fill=WHITE, radius=0.16, alpha=0.09)
    val_col = RGBColor(0x9C, 0xE8, 0xB0) if val == "LIVE" else RGBColor(0x8F, 0xE0, 0xCF)
    tb(s, chx, 5.0, w, 0.45, val, 19, val_col, True, align=PP_ALIGN.CENTER, font=FS)
    tb(s, chx, 5.44, w, 0.3, lab, 10, OFFW, align=PP_ALIGN.CENTER)
    chx += w + 0.22

tb(s, 0.75, 6.75, 10.5, 0.3, "PROJECT ai-ippc   ·   REGION us-central1   ·   Prepared from the deployed GCP environment",
   9.5, RGBColor(0x7E, 0x95, 0xA8), spc=80)
notes(s, "Takeaway: We turned the IPCC's massive, authoritative climate reports into a live assistant on Google Cloud that answers in minutes — and every answer is grounded in the original evidence.")

# ================================================================ SLIDE 2 — WHY (before / after)
s = slide_new()
header(s, "Why this matters", "Authoritative evidence, locked in thousands of pages",
       "The assistant closes the gap between rigorous IPCC science and the people who need it daily.", 2)

# BEFORE panel
rrect(s, 0.6, 2.0, 5.35, 3.35, fill=PANEL, radius=0.06)
tb(s, 0.95, 2.25, 4.7, 0.3, "TODAY — MANUAL RESEARCH", 10.5, GRAY, True, spc=140)
before = ["Weeks of searching across chapters and working groups",
          "Deep domain expertise needed just to locate a passage",
          "Answers are hard to trace back to a defensible source"]
for i, t in enumerate(before):
    y = 2.75 + i * 0.82
    dot(s, 0.95, y, 0.3, RGBColor(0xF2, 0xDE, 0xDE))
    tb(s, 0.95, y - 0.005, 0.3, 0.3, "✕", 12, RED, True, align=PP_ALIGN.CENTER)
    tb(s, 1.42, y - 0.04, 4.25, 0.7, t, 11.5, INK, line_sp=1.1)

# arrow
ar = shp(s, MSO_SHAPE.RIGHT_ARROW, 6.12, 3.4, 1.05, 0.55, fill=TEAL)
try:
    ar.adjustments[0] = 0.55; ar.adjustments[1] = 0.55
except Exception:
    pass

# AFTER panel
rrect(s, 7.35, 2.0, 5.35, 3.35, fill=TEALBG, line=GREEN, lw=1.0, radius=0.06)
tb(s, 7.7, 2.25, 4.7, 0.3, "WITH THE ASSISTANT — EVIDENCE ON DEMAND", 10.5, RGBColor(0x1E, 0x7A, 0x52), True, spc=120)
after = ["Minutes to a clear answer written in plain language",
         "Anyone can ask — planners, students, policy teams",
         "Every answer cites its passages and links the original PDF"]
for i, t in enumerate(after):
    y = 2.75 + i * 0.82
    dot(s, 7.7, y, 0.3, RGBColor(0xCF, 0xEA, 0xDC))
    tb(s, 7.7, y - 0.01, 0.3, 0.3, "✓", 12, GREEN, True, align=PP_ALIGN.CENTER)
    tb(s, 8.17, y - 0.04, 4.25, 0.7, t, 11.5, INK, line_sp=1.1)

# operating principle strip
tb(s, 0.6, 5.72, 2.6, 0.3, "OPERATING PRINCIPLE", 10, FAINT, True, spc=140)
pills = [("1", "Retrieve evidence", NAVY), ("2", "Generate explanation", TEAL), ("3", "Show sources", GREEN)]
px = 3.35
for i, (n, t, col) in enumerate(pills):
    w = 2.75
    rrect(s, px, 5.62, w, 0.55, fill=col, radius=0.5)
    tb(s, px, 5.74, w, 0.32, t, 12, WHITE, True, align=PP_ALIGN.CENTER)
    if i < 2:
        tb(s, px + w + 0.02, 5.66, 0.3, 0.4, "→", 15, FAINT, True, align=PP_ALIGN.CENTER)
    px += w + 0.34
footer(s, "Audience: climate teams, city planners, students and policy stakeholders — with source visibility and original PDF access.")
notes(s, "Takeaway: The problem is access, not evidence. We move from weeks of expert searching to minutes of grounded conversation — and rigor survives because every answer shows its sources.")

# ================================================================ SLIDE 3 — IPCC ASSESSMENT LIFECYCLE (serpentine journey)
s = slide_new()
header(s, "Context — how the source reports are made", "The five-to-seven-year lifecycle behind every IPCC report",
       "Each report in the corpus survives scoping, drafting, two review rounds and line-by-line government approval.", 3, accent=BLUE)

TRACK = RGBColor(0xD5, 0xDD, 0xE3)
steps10 = [
    ("1",  "Scope",             "Governments approve outline,\nschedule and budget", NAVY),
    ("2",  "Select authors",    "Bureau picks balanced\nexpert teams",               NAVY),
    ("3",  "Assess evidence",   "Published literature only —\nno new research",      TEAL),
    ("4",  "First draft",       "Three Working Groups\ndraft chapters",              TEAL),
    ("5",  "Expert review",     "Every comment answered\non the record",             AMBER),
    ("6",  "Government review", "Second draft + Summary\nfor Policymakers",          AMBER),
    ("7",  "Final draft",       "Authors respond; final\nSPM circulated",            AMBER),
    ("8",  "Endorsement",       "Plenary approves the SPM\nline by line",            GREEN),
    ("9",  "Publication",       "Report, comments and\nresponses released",          GREEN),
    ("10", "Synthesis",         "Working Groups integrated;\nnext cycle begins",     PURPLE),
]
xs1 = [1.7, 4.0, 6.3, 8.6, 10.9]
y1, y2 = 3.05, 5.45
xs2 = [10.9, 8.6, 6.3, 4.0, 1.7]

# serpentine track: row 1 → right bend → row 2 → next-cycle terminal
connector(s, 1.7, y1, 11.55, y1, color=TRACK, w=2.5)
connector(s, 11.55, y1, 11.55, y2, color=TRACK, w=2.5)
connector(s, 11.55, y2, 1.35, y2, color=TRACK, w=2.5, arrow=True)
dot(s, 11.55 - 0.05, y1 - 0.05, 0.1, TRACK)
dot(s, 11.55 - 0.05, y2 - 0.05, 0.1, TRACK)

for idx, (num, ttl, desc, col) in enumerate(steps10):
    if idx < 5:
        x, yy = xs1[idx], y1
        ty, dy = yy - 0.78, yy + 0.38
    else:
        x, yy = xs2[idx - 5], y2
        ty, dy = yy - 0.78, yy + 0.38
    r = 0.27
    dot(s, x - r, yy - r, 2*r, WHITE, line=col, lw=2.75)
    tb(s, x - r, yy - 0.155, 2*r, 0.3, num, 12, col, True, align=PP_ALIGN.CENTER, font=FS)
    tb(s, x - 1.1, ty, 2.2, 0.32, ttl, 12, INK, True, align=PP_ALIGN.CENTER, font=FS)
    tb(s, x - 1.1, dy, 2.2, 0.55, desc, 8.5, GRAY, align=PP_ALIGN.CENTER, line_sp=1.05)

# next-cycle terminal (icon only — step 10's caption already says "next cycle begins")
loop_icon = shp(s, MSO_SHAPE.CIRCULAR_ARROW, 0.62, y2 - 0.26, 0.52, 0.52, fill=PURPLE)

# phase legend
phases = [("DEFINE", NAVY), ("ASSESS & DRAFT", TEAL), ("REVIEW & REFINE", AMBER),
          ("APPROVE & PUBLISH", GREEN), ("SYNTHESIZE", PURPLE)]
lx = 2.05
for name, col in phases:
    dot(s, lx, 6.55, 0.14, col)
    w_lab = 0.22 + 0.093 * len(name)
    tb(s, lx + 0.22, 6.5, w_lab, 0.25, name, 8.5, GRAY, True, spc=60)
    lx += w_lab + 0.55

footer(s, "Sources: IPCC — Preparing Reports (ipcc.ch/about/preparingreports) · IPCC procedures (ipcc.ch/documentation/procedures) · IPCC Review Process factsheet, 2024", accent=BLUE)
notes(s, "Takeaway: The corpus is not just literature — every report survived a 5–7 year process with two review rounds, on-the-record author responses, and line-by-line government approval of the Summary for Policymakers. That is why grounding answers in it makes them defensible.")

# ================================================================ SLIDE 4 — WHY THE IPCC NEEDS AI (SIGIR / MANILA24)
s = slide_new()
header(s, "The scale problem", "Why the IPCC needs AI: evidence grows faster than any author team",
       "One working-group report already assesses tens of thousands of papers — and the next cycle (AR7) faces more volume and more blind spots.", 4, accent=AMBER, hsize=24)

# stat bubble cluster (left) — one AR6 working-group report
bubbles = [
    (2.35, 3.60, 1.05, NAVY,  "34,000+", "papers assessed", 20, 9.5),
    (4.42, 2.78, 0.62, TEAL,  "270",     "authors",         16, 8.5),
    (4.62, 4.38, 0.55, GREEN, "67",      "countries",       15, 8.5),
    (3.18, 5.28, 0.45, AMBER, "18",      "chapters",        14, 8),
]
for bx, by, br, col, val, lab, vs, ls in bubbles:
    dot(s, bx - br, by - br, 2*br, col)
    tb(s, bx - br, by - 0.30, 2*br, 0.4, val, vs, WHITE, True, align=PP_ALIGN.CENTER, font=FS)
    tb(s, bx - br, by + 0.06, 2*br, 0.3, lab, ls, RGBColor(0xDC, 0xE8, 0xEF), align=PP_ALIGN.CENTER)
tb(s, 0.6, 5.95, 5.2, 0.6, "One AR6 working-group report alone\n(WGII — impacts, adaptation and vulnerability)", 9.5, FAINT, align=PP_ALIGN.CENTER, line_sp=1.15)

# AR7 twin challenge (right, top)
rrect(s, 5.95, 2.0, 6.78, 1.62, fill=AMBERBG, radius=0.08)
tb(s, 6.25, 2.18, 6.2, 0.28, "THE AR7 TWIN CHALLENGE", 10, RGBColor(0xA5, 0x6E, 0x1B), True, spc=140)
multi(s, 6.25, 2.55, 3.05, 0.95,
      [[("More volume", FS, 11.5, INK, True)],
       [("Climate literature is growing faster than author teams can sift.", F, 9.5, GRAY, False)]], line_sp=1.1, space_after=3)
multi(s, 9.55, 2.55, 2.95, 0.95,
      [[("More blind spots", FS, 11.5, INK, True)],
       [("Key regions and topics have little published English literature.", F, 9.5, GRAY, False)]], line_sp=1.1, space_after=3)

# SIGIR research agenda (right, bottom)
rrect(s, 5.95, 3.85, 6.78, 2.5, fill=WHITE, line=RGBColor(0xDD, 0xE4, 0xE9), lw=1.0, radius=0.08)
tb(s, 6.25, 4.03, 6.3, 0.28, "WHAT AI MUST DELIVER — SIGIR / MANILA24 RESEARCH AGENDA", 10, FAINT, True, spc=110)
agenda = [("Transparent, traceable evidence synthesis — every claim carries its source", TEAL),
          ("Multilingual sources and grey literature widen the evidence net", GREEN),
          ("Bias mitigation across geographies, cultures and data-poor regions", AMBER),
          ("Multi-modal retrieval: text, data, maps and community knowledge", PURPLE)]
for i, (t, col) in enumerate(agenda):
    yy = 4.42 + i * 0.45
    dot(s, 6.28, yy + 0.03, 0.16, col)
    tb(s, 6.58, yy - 0.04, 6.0, 0.35, t, 10.5, INK)
footer(s, "Already live in this platform: retrieval-first answers with citations.   Source: van den Hurk, de Rijke & Salim (eds.), “Information Retrieval for Climate Impact” (MANILA24), ACM SIGIR Forum 59(1), June 2025.", accent=AMBER)
notes(s, "Takeaway: This is not AI for AI's sake. AR6 WGII alone synthesized 34,000+ papers with 270 authors from 67 countries, and AR7 faces more literature and more blind spots. The information-retrieval research community (SIGIR MANILA24 workshop) has set an agenda — traceability, multilingual and grey literature, bias mitigation — and our platform already implements its first principle: retrieval with citations.")

# ================================================================ SLIDE 5 — WHY CITIES (three realities)
s = slide_new()
header(s, "Who it serves", "Cities feel climate change first — and decide under pressure",
       "Urban climate decisions are interconnected, time-sensitive and highly visible.", 5, accent=RED)

realities = [
    ("01", "The city reality", RED,
     "Heat, flooding, water stress, infrastructure disruption and health impacts compound each other — and exposure is unequal."),
    ("02", "The information reality", AMBER,
     "Relevant evidence is scattered across long reports, chapters, annexes and technical language. Search time becomes decision delay."),
    ("03", "The leadership reality", TEAL,
     "Executives need concise implications, practitioners need detail — and everyone needs a route back to the original source."),
]
for i, (num, ttl, col, txt) in enumerate(realities):
    x = 0.6 + i * 4.12
    rrect(s, x, 2.05, 3.9, 2.85, fill=WHITE, line=RGBColor(0xDD, 0xE4, 0xE9), lw=1.0, radius=0.07)
    shp(s, MSO_SHAPE.RECTANGLE, x, 2.05, 3.9, 0.09, fill=col)
    tb(s, x + 0.3, 2.35, 1.0, 0.5, num, 22, col, True, font=FS)
    tb(s, x + 0.3, 2.95, 3.3, 0.35, ttl, 14, INK, True, font=FS)
    tb(s, x + 0.3, 3.4, 3.3, 1.35, txt, 10.5, GRAY, line_sp=1.2)

tb(s, 0.6, 5.58, 2.55, 0.6, "STRATEGIC\nOPPORTUNITY", 10, FAINT, True, spc=140, line_sp=1.2)
chev = [("Question", NAVY), ("Evidence", TEAL), ("Informed action", GREEN)]
cxx = 3.3
for t, col in chev:
    c = shp(s, MSO_SHAPE.CHEVRON, cxx, 5.5, 3.0, 0.62, fill=col)
    try: c.adjustments[0] = 0.5
    except Exception: pass
    tb(s, cxx + 0.25, 5.65, 2.6, 0.32, t, 12, WHITE, True, align=PP_ALIGN.CENTER)
    cxx += 3.12
footer(s, "Shorten the path from question to evidence to action — for planners, operators and leadership alike.", accent=RED)
notes(s, "Takeaway: Cities are where climate impacts and accountability meet. The bottleneck is not evidence — it is the time between a question and a defensible answer. That path is what we compress.")

# ================================================================ SLIDE 6 — THE EXECUTIVE CASE (three lenses)
s = slide_new()
header(s, "The executive case", "Speed, control and trust — one platform, three lenses",
       "A governed evidence assistant connects outcomes for the CEO, control for the CIO and confidence for risk.", 6)

lenses = [
    ("OUTCOME", "CEO lens", GREEN,
     ["Cut research cycle time from weeks to minutes", "Consistent, source-backed briefings",
      "Wider access to authoritative knowledge", "Better-supported resilience decisions"]),
    ("CONTROL", "CIO lens", NAVY,
     ["Answers grounded in approved sources only", "Full traceability from claim to passage",
      "Storage, retrieval and generation separated", "Identity, cost and operations in hand"]),
    ("TRUST", "Risk lens", AMBER,
     ["Uncertainty made visible, not hidden", "No unsupported claims or stale knowledge",
      "No uncontrolled access or model overreach", "Human judgment stays accountable"]),
]
for i, (tag, lens, col, items) in enumerate(lenses):
    x = 0.6 + i * 4.12
    rrect(s, x, 2.05, 3.9, 3.55, fill=WHITE, line=RGBColor(0xDD, 0xE4, 0xE9), lw=1.0, radius=0.07)
    shp(s, MSO_SHAPE.RECTANGLE, x, 2.05, 3.9, 0.52, fill=col)
    multi(s, x + 0.3, 2.16, 3.4, 0.35,
          [[(tag + "   ", FS, 12, WHITE, True, 80), ("— " + lens, F, 10.5, RGBColor(0xE6, 0xEE, 0xF3), False)]])
    for j, it in enumerate(items):
        yy = 2.85 + j * 0.66
        dot(s, x + 0.3, yy + 0.04, 0.13, col)
        tb(s, x + 0.56, yy - 0.04, 3.1, 0.6, it, 10, GRAY, line_sp=1.1)
rrect(s, 2.6, 5.85, 8.15, 0.58, fill=PANEL, radius=0.5)
tb(s, 2.6, 5.99, 8.15, 0.32, "Faster insight is valuable only when it is governed, explainable and sustainable.", 11.5, INK, True, align=PP_ALIGN.CENTER)
footer(s, "One platform serves all three lenses — the same retrieval-first design delivers speed, traceability and control.")
notes(s, "Takeaway: The CEO gets speed and consistency, the CIO gets traceability and control, risk gets visible uncertainty and human accountability — from one design decision: retrieve approved evidence before generating a word.")

# ================================================================ SLIDE 7 — WHAT (hub & spoke)
s = slide_new()
header(s, "What it does", "One assistant, five ways to work with climate reports",
       "Every workflow runs through the same conversation — no tooling change, no training curve.", 7)

hub_x, hub_y = 6.667, 4.5
sat = [("Ask",       "Semantic search with\nsourced answers",      TEAL,   90),
       ("Browse",    "Open or download\nthe original PDFs",        BLUE,   162),
       ("Summarize", "Report-level summaries\non demand",          GREEN,  234),
       ("Review",    "Critique and assess\ndocuments",             AMBER,  306),
       ("Improve",   "Suggest clearer,\nstronger content",         PURPLE, 18)]
R = 1.85
for name, desc, col, ang in sat:
    a = math.radians(ang)
    sx = hub_x + R * math.cos(a)
    sy = hub_y - R * math.sin(a) * 0.82   # slight vertical squash
    connector(s, hub_x, hub_y, sx, sy, color=RGBColor(0xC9, 0xD4, 0xDC), w=1.2)

for name, desc, col, ang in sat:
    a = math.radians(ang)
    sx = hub_x + R * math.cos(a)
    sy = hub_y - R * math.sin(a) * 0.82
    r = 0.62
    dot(s, sx - r, sy - r, 2*r, col)
    tb(s, sx - r, sy - 0.16, 2*r, 0.35, name, 12.5, WHITE, True, align=PP_ALIGN.CENTER, font=FS)
    # description further out along the spoke (top satellite: to the right of its circle)
    if ang == 90:
        tb(s, sx + 0.78, sy - 0.28, 2.4, 0.6, desc, 10, GRAY, align=PP_ALIGN.LEFT, line_sp=1.1)
    else:
        dx = hub_x + 3.35 * math.cos(a)
        dy = hub_y - 3.35 * math.sin(a) * 0.78
        tb(s, dx - 1.35, dy - 0.3, 2.7, 0.7, desc, 10, GRAY, align=PP_ALIGN.CENTER, line_sp=1.1)

hr = 0.95
dot(s, hub_x - hr, hub_y - hr, 2*hr, NAVY)
dot(s, hub_x - hr - 0.09, hub_y - hr - 0.09, 2*hr + 0.18, None, line=NAVY, lw=1.0)
tb(s, hub_x - hr, hub_y - 0.36, 2*hr, 0.5, "IPCC AI", 15, WHITE, True, align=PP_ALIGN.CENTER, font=FS)
tb(s, hub_x - hr, hub_y + 0.02, 2*hr, 0.3, "Assistant", 11, RGBColor(0xA8, 0xC4, 0xD8), align=PP_ALIGN.CENTER)

footer(s, "Research journey: Discover → Understand → Verify → Communicate   ·   Python services on Cloud Run   ·   Intelligence: Vertex AI")
notes(s, "Takeaway: This is not a search box — it is five research workflows in one conversation: ask, browse, summarize, review and improve, all against the same trusted corpus.")

# ================================================================ SLIDE 8 — ARCHITECTURE (layered landscape)
s = slide_new()
header(s, "How it is built", "A three-layer platform: knowledge, intelligence, experience",
       "Managed GCP services separate storage, retrieval, reasoning and presentation — nothing to patch, everything scales.", 8)

def layer(y, h, tint, label, label_col):
    rrect(s, 1.95, y, 9.55, h, fill=tint, radius=0.07)
    tb(s, 2.2, y + 0.13, 7.0, 0.28, label, 10, label_col, True, spc=160)

def pill(x, y, w, name, detail, col):
    rrect(s, x, y, w, 0.78, fill=WHITE, line=RGBColor(0xDD, 0xE4, 0xE9), lw=0.75, radius=0.18)
    dot(s, x + 0.16, y + 0.245, 0.28, col)
    multi(s, x + 0.56, y + 0.115, w - 0.68, 0.62,
          [[(name, FS, 11.5, INK, True)], [(detail, F, 9, GRAY, False)]], line_sp=1.0)

# experience (top)
layer(1.98, 1.22, NAVYBG, "EXPERIENCE  —  WHERE PEOPLE WORK", NAVY)
pill(4.35, 2.3, 4.7, "Cloud Run  ·  Streamlit assistant", "Public endpoint · 1 warm instance · 100% traffic", NAVY)

# intelligence (middle)
layer(3.42, 1.5, TEALBG, "INTELLIGENCE  —  WHERE MEANING IS MADE", RGBColor(0x0A, 0x63, 0x76))
pill(2.5, 3.9, 2.85, "Vertex AI Embeddings", "text-embedding-005", TEAL)
pill(5.55, 3.9, 2.85, "BigQuery Vector Search", "top-k semantic retrieval", TEAL)
pill(8.6, 3.9, 2.7, "Gemini 3 Flash", "grounded generation (preview)", GREEN)

# knowledge (bottom)
layer(5.14, 1.35, AMBERBG, "KNOWLEDGE  —  WHERE EVIDENCE LIVES", RGBColor(0xA5, 0x6E, 0x1B))
pill(3.3, 5.55, 3.3, "Cloud Storage", "8 IPCC PDF reports · ≈73 MiB", AMBER)
pill(7.0, 5.55, 3.3, "BigQuery documents", "10,184 extracted passages", AMBER)

# left rail: offline knowledge pipeline (up)
connector(s, 1.45, 6.35, 1.45, 2.15, color=AMBER, w=2.0, arrow=True)
lbl = tb(s, -0.62, 4.05, 4.1, 0.3, "KNOWLEDGE PIPELINE (OFFLINE) · UPLOAD → CHUNK → EMBED → INDEX", 8.5, RGBColor(0xA5, 0x6E, 0x1B), True, align=PP_ALIGN.CENTER, spc=60, wrap=False)
lbl.rotation = 270

# right rail: live question journey (down + up)
connector(s, 11.95, 2.15, 11.95, 6.35, color=TEAL, w=2.0, arrow=True, both=True)
lbl2 = tb(s, 10.35, 4.05, 4.1, 0.3, "QUESTION JOURNEY (LIVE) · ASK → RETRIEVE → ANSWER", 8.5, RGBColor(0x0A, 0x63, 0x76), True, align=PP_ALIGN.CENTER, spc=60, wrap=False)
lbl2.rotation = 90

footer(s, "Offline pipeline builds the knowledge once; the live path reuses it for every question at low cost.")
notes(s, "Takeaway: Read the slide bottom-up — evidence lives in the knowledge layer, meaning is made in the intelligence layer, people work in the experience layer. Amber arrow = offline ingestion; teal arrow = every live question.")

# ================================================================ SLIDE 9 — RAG JOURNEY (metro line)
s = slide_new()
header(s, "How answers are made", "The journey from question to defensible answer",
       "Five stations — retrieval comes before generation, so the model explains evidence instead of inventing it.", 9)

line_y = 4.15
stations = [("1", "Ask",      "A user poses a climate\nquestion in plain words", RGBColor(0x0B, 0x2E, 0x4F)),
            ("2", "Embed",    "Vertex AI turns the question\ninto a semantic vector",  RGBColor(0x14, 0x53, 0x71)),
            ("3", "Retrieve", "BigQuery returns the closest\nIPCC passages (top-k)",   RGBColor(0x0C, 0x7B, 0x93)),
            ("4", "Generate", "Gemini writes the answer\nfrom those passages only",    RGBColor(0x1F, 0x8D, 0x80)),
            ("5", "Cite",     "The UI shows the answer\nwith sources and PDF links",   GREEN)]
xs = [1.35, 3.55, 5.75, 7.95, 10.15]
# track segments
for i in range(len(xs) - 1):
    connector(s, xs[i], line_y, xs[i+1], line_y, color=stations[i+1][3], w=5.0)
connector(s, xs[-1], line_y, 11.6, line_y, color=GREEN, w=5.0)
# terminal pill
rrect(s, 11.35, 3.82, 1.55, 0.66, fill=GREEN, radius=0.5)
tb(s, 11.35, 3.9, 1.55, 0.55, "Grounded,\ncited answer", 9.5, WHITE, True, align=PP_ALIGN.CENTER, line_sp=1.0)

for (num, name, desc, col), x in zip(stations, xs):
    r = 0.31
    dot(s, x - r, line_y - r, 2*r, WHITE, line=col, lw=3.0)
    tb(s, x - r, line_y - 0.17, 2*r, 0.35, num, 14, col, True, align=PP_ALIGN.CENTER, font=FS)
    tb(s, x - 1.0, line_y - 0.95, 2.0, 0.35, name, 14, INK, True, align=PP_ALIGN.CENTER, font=FS)
    tb(s, x - 1.05, line_y + 0.5, 2.1, 0.75, desc, 9.5, GRAY, align=PP_ALIGN.CENTER, line_sp=1.1)

# why it matters panel
rrect(s, 2.3, 5.9, 8.7, 0.72, fill=TEALBG, radius=0.14)
multi(s, 2.65, 6.05, 8.1, 0.5,
      [[("Why it matters   ", FS, 11, RGBColor(0x0A, 0x63, 0x76), True),
        ("Retrieval grounds every response in IPCC content and sharply reduces unsupported generation.", F, 11, INK, False)]],
      anchor=MSO_ANCHOR.MIDDLE)
footer(s, "Example: “How can cities reduce the risks of extreme heat for vulnerable residents?” — answer + passages + original PDF = a verifiable loop.")
notes(s, "Takeaway: Walk the line left to right — the key design choice is station 3 before station 4. Evidence is fetched before a single word is generated, which is what makes answers defensible.")

# ================================================================ SLIDE 10 — PROOF (KPI dashboard)
s = slide_new()
header(s, "Where it stands", "Live and verified in project ai-ippc",
       "Every figure below was read from the deployed environment — not from a plan.", 10, accent=GREEN)

# LIVE badge
rrect(s, 10.05, 0.62, 2.68, 0.52, fill=RGBColor(0xDD, 0xF2, 0xE4), radius=0.5)
dot(s, 10.28, 0.79, 0.18, GREEN)
tb(s, 10.55, 0.72, 2.1, 0.3, "LIVE — HTTP 200 VERIFIED", 9.5, RGBColor(0x1E, 0x7A, 0x52), True, spc=60)

tiles = [
    ("SERVICE", "100%", "traffic on latest revision", ["Cloud Run · ipcc-ai-project", "Revision 00007-lfs", "1 always-warm instance"], NAVY),
    ("CORPUS", "8", "IPCC reports online", ["Cloud Storage", "ipcc-srcities-bucket-1", "≈73 MiB of source PDFs"], AMBER),
    ("INDEX", "10,184", "passages embedded", ["BigQuery · climate_ai", "documents = embeddings", "full corpus coverage"], TEAL),
    ("MODELS", "2", "Vertex AI models", ["Gemini 3 Flash (preview)", "text-embedding-005", "Region us-central1"], GREEN),
]
tx = 0.6
for cap, big, unit, subs, col in tiles:
    w = 3.0
    rrect(s, tx, 2.05, w, 3.15, fill=WHITE, line=RGBColor(0xDD, 0xE4, 0xE9), lw=1.0, radius=0.07)
    shp(s, MSO_SHAPE.RECTANGLE, tx, 2.05, w, 0.09, fill=col)
    tb(s, tx + 0.3, 2.35, w - 0.6, 0.3, cap, 10, FAINT, True, spc=160)
    tb(s, tx + 0.3, 2.68, w - 0.6, 0.75, big, 34, col, True, font=FS)
    tb(s, tx + 0.3, 3.42, w - 0.6, 0.3, unit, 11, INK, True)
    for i, sub in enumerate(subs):
        tb(s, tx + 0.3, 3.85 + i * 0.34, w - 0.6, 0.3, sub, 9.5, GRAY)
    tx += w + 0.19

# coverage bar
tb(s, 0.6, 5.6, 8.0, 0.3, "INDEX COVERAGE — EVERY STORED PASSAGE IS EMBEDDED", 9.5, FAINT, True, spc=120)
rrect(s, 0.6, 5.95, 9.6, 0.34, fill=RGBColor(0xE4, 0xEA, 0xEF), radius=0.5)
rrect(s, 0.6, 5.95, 9.6, 0.34, fill=GREEN, radius=0.5)
tb(s, 10.4, 5.93, 2.3, 0.35, "10,184 / 10,184 · 100%", 12, GREEN, True)
footer(s, "Project number 510679909889   ·   Billing enabled   ·   Region us-central1", accent=GREEN)
notes(s, "Takeaway: This is a verified deployment, not a proposal. The service is serving 100% of traffic, the full corpus is indexed 1:1, and the endpoint returns HTTP 200 publicly.")

# ================================================================ SLIDE 11 — GOVERNANCE (six controls)
s = slide_new()
header(s, "Governance", "Trust is not a feature — it is the operating model around AI",
       "Six controls a CIO can evidence, each mapped to the platform as it scales.", 11, accent=NAVY)

controls = [
    ("Accuracy", "Faithfulness and source-relevance evaluation on golden questions", TEAL),
    ("Security", "Least-privilege identity and controlled access to service and data", NAVY),
    ("Freshness", "Versioned documents and automated ingestion keep knowledge current", GREEN),
    ("Operations", "Health checks, logs, alerts and quota monitoring end to end", AMBER),
    ("Economics", "Budgets, rate limits and cost dashboards for AI and data services", BLUE),
    ("Accountability", "Human review stays in the loop for high-impact decisions", PURPLE),
]
for i, (ttl, txt, col) in enumerate(controls):
    x = 0.6 + (i % 3) * 4.12
    y = 2.05 + (i // 3) * 2.15
    rrect(s, x, y, 3.9, 1.95, fill=WHITE, line=RGBColor(0xDD, 0xE4, 0xE9), lw=1.0, radius=0.09)
    dot(s, x + 0.3, y + 0.28, 0.46, None, line=col, lw=2.25)
    tb(s, x + 0.3, y + 0.36, 0.46, 0.3, "✓", 13, col, True, align=PP_ALIGN.CENTER)
    tb(s, x + 0.95, y + 0.3, 2.75, 0.32, ttl, 13.5, INK, True, font=FS)
    tb(s, x + 0.95, y + 0.72, 2.75, 1.1, txt, 10, GRAY, line_sp=1.18)
footer(s, "A CEO needs outcomes; a CIO needs evidence that the outcomes are controlled.", accent=NAVY)
notes(s, "Takeaway: These six controls are the production gate. None of them require new technology — they are configuration, evaluation and process around the platform that already runs.")

# ================================================================ SLIDE 12 — RESILIENCE (incident timeline)
s = slide_new()
header(s, "What we learned", "One incident, resolved — and the platform is stronger for it",
       "A launch-day outage was diagnosed in logs, fixed at the root cause, and verified from the outside.", 12, accent=AMBER)

ty = 2.85
connector(s, 1.3, ty, 12.0, ty, color=RGBColor(0xD5, 0xDD, 0xE3), w=3.0)
phases = [
    (2.35, RED,   "✕", "INCIDENT", "Service down",
     "Cloud Run returned 500/503 errors. Cloud Logging traced the root cause: billing was disabled at project level."),
    (6.65, AMBER, "!", "RESPONSE", "Root cause fixed",
     "Billing account attached and the service redeployed — memory raised to 1 GiB with one always-warm instance."),
    (10.95, GREEN, "✓", "RESOLVED", "Live and verified",
     "The public endpoint answers HTTP 200 and is ready for functional testing of all five workflows."),
]
for x, col, mark, tag, ttl, txt in phases:
    dot(s, x - 0.27, ty - 0.27, 0.54, col)
    tb(s, x - 0.27, ty - 0.16, 0.54, 0.3, mark, 14, WHITE, True, align=PP_ALIGN.CENTER)
    tb(s, x - 1.85, ty - 0.75, 3.7, 0.3, tag, 10, col, True, align=PP_ALIGN.CENTER, spc=160)
    rrect(s, x - 1.85, ty + 0.55, 3.7, 1.95, fill=WHITE, line=RGBColor(0xDD, 0xE4, 0xE9), lw=1.0, radius=0.09)
    shp(s, MSO_SHAPE.RECTANGLE, x - 1.85, ty + 0.55, 0.07, 1.95, fill=col)
    tb(s, x - 1.55, ty + 0.75, 3.2, 0.3, ttl, 13, INK, True, font=FS)
    tb(s, x - 1.55, ty + 1.12, 3.25, 1.3, txt, 10, GRAY, line_sp=1.15)

# lessons strip
tb(s, 0.6, 5.85, 2.9, 0.3, "WHAT WE KEEP", 10, FAINT, True, spc=140)
lessons = [("Logs find root causes fast", TEAL), ("Keep capacity headroom", AMBER), ("Verify from the outside in", GREEN)]
lx = 2.6
for t, col in lessons:
    w = 3.15
    rrect(s, lx, 5.75, w, 0.5, fill=WHITE, line=col, lw=1.0, radius=0.5)
    tb(s, lx, 5.86, w, 0.3, t, 10.5, INK, True, align=PP_ALIGN.CENTER)
    lx += w + 0.3
footer(s, "Operational next step: exercise every UI workflow and monitor Vertex AI / BigQuery quotas.", accent=AMBER)
notes(s, "Takeaway: The only outage was environmental (billing), not architectural. Diagnosis took minutes because logging was in place — and the fix left the service with more headroom than before.")

# ================================================================ SLIDE 13 — ROADMAP (staircase)
s = slide_new()
header(s, "Where we go next", "From working deployment to production climate intelligence",
       "Each step compounds the last — validate what exists, then harden, secure, operate and scale it.", 13, accent=GREEN, hsize=25)

steps = [
    ("Validate", "UAT across all five workflows; verify source accuracy.", RGBColor(0x0B, 0x2E, 0x4F)),
    ("Harden",   "Config via environment variables; health checks; retire stale scripts.", RGBColor(0x14, 0x53, 0x71)),
    ("Secure",   "Least-privilege service account; restrict Cloud Run access.", RGBColor(0x0C, 0x7B, 0x93)),
    ("Operate",  "Logging, alerting, quota and cost monitoring.", RGBColor(0x1F, 0x8D, 0x80)),
    ("Scale",    "Automated document ingestion; role-based access for teams.", GREEN),
]
base_y = 6.35
for i, (ttl, desc, col) in enumerate(steps):
    x = 0.75 + i * 2.42
    w = 2.22
    top = 5.45 - i * 0.72
    rrect(s, x, top, w, base_y - top, fill=col, radius=0.05)
    tb(s, x, top + 0.12, w, 0.35, ttl, 14, WHITE, True, align=PP_ALIGN.CENTER, font=FS)
    if i < 4:
        tb(s, x, top - 0.98, w, 0.9, desc, 9.5, GRAY, align=PP_ALIGN.CENTER, line_sp=1.12)
    else:
        # tallest bar: caption fits inside, freeing the summit area above
        tb(s, x + 0.18, top + 0.62, w - 0.36, 0.9, desc, 9.5, RGBColor(0xE0, 0xF3, 0xE9), align=PP_ALIGN.CENTER, line_sp=1.15)

# phase labels
connector(s, 0.75, base_y + 0.12, 12.61, base_y + 0.12, color=RGBColor(0xD5, 0xDD, 0xE3), w=1.0)
for label, x, w in (("NOW", 0.75, 4.64), ("NEXT", 5.59, 4.64), ("LATER", 10.43, 2.22)):
    tb(s, x, base_y + 0.2, w, 0.3, label, 9.5, FAINT, True, align=PP_ALIGN.CENTER, spc=200)

# summit marker above the Scale bar (its caption now lives inside the bar)
flag = shp(s, MSO_SHAPE.ISOSCELES_TRIANGLE, 10.48, 1.82, 0.3, 0.24, fill=GREEN)
flag.rotation = 90
tb(s, 10.88, 1.72, 1.8, 0.55, "PRODUCTION-READY\nCLIMATE INTELLIGENCE", 9, RGBColor(0x1E, 0x7A, 0x52), True, align=PP_ALIGN.LEFT, line_sp=1.15, spc=80)

footer(s, "Decision ask: sponsor a measured pilot with explicit quality, security and cost gates.", accent=GREEN)
notes(s, "Takeaway: We're asking for a decision to move along this staircase: validate and harden now, secure and operate next, then scale to more documents and more teams.")

# ================================================================ SLIDE 14 — FUTURE STATE (Enterprise Knowledge Fabric)
s = slide_new()
header(s, "Future state architecture", "An Enterprise Knowledge Fabric with reasoning agents",
       "The platform inherits the seven-layer fabric — a new ontology layer and enterprise context graph give agents shared meaning to reason over.",
       14, accent=PURPLE, hsize=25)

# seven layers, stacked bottom-up (L1 at bottom); L4/L5 are new, L6 evolves into agents
LX, LW, LH, LGAP = 0.95, 8.3, 0.56, 0.06
layers7 = [   # (level, name, detail, style)
    ("L7", "Experience & Governance",          "workspaces · APIs · identity · observability · cost",              "base"),
    ("L6", "Consumption & Reasoning — Agents", "research, synthesis, review and governance agents (Gemini)",       "agents"),
    ("L5", "Enterprise Context Graph",         "entities, relationships and lineage linked across every source",   "new-purple"),
    ("L4", "Ontology Layer",                   "shared climate vocabulary: hazards, sectors, regions, measures",   "new-green"),
    ("L3", "Semantic Enrichment",              "chunking · embeddings · entity extraction (Vertex AI)",            "today"),
    ("L2", "Storage & Processing",             "Cloud Storage · BigQuery · ingestion pipelines",                   "today"),
    ("L1", "Sources & Ingestion",              "IPCC reports · grey literature · city data · multilingual corpora","today"),
]
ys = [2.20 + i * (LH + LGAP) for i in range(7)]
for (lvl, name, detail, style), y in zip(layers7, ys):
    if style == "agents":
        rrect(s, LX, y, LW, LH, fill=NAVY, radius=0.1)
        name_col, det_col = WHITE, RGBColor(0xA8, 0xC4, 0xD8)
        badge = ("EVOLVED", AMBER)
    elif style == "new-green":
        rrect(s, LX, y, LW, LH, fill=WHITE, line=GREEN, lw=1.75, radius=0.1)
        name_col, det_col = INK, GRAY
        badge = ("NEW", GREEN)
    elif style == "new-purple":
        rrect(s, LX, y, LW, LH, fill=WHITE, line=PURPLE, lw=1.75, radius=0.1)
        name_col, det_col = INK, GRAY
        badge = ("NEW", PURPLE)
    elif style == "base":
        rrect(s, LX, y, LW, LH, fill=NAVYBG, radius=0.1)
        name_col, det_col = NAVY, GRAY
        badge = None
    else:
        rrect(s, LX, y, LW, LH, fill=PANEL, radius=0.1)
        name_col, det_col = INK, GRAY
        badge = ("TODAY", FAINT)
    tb(s, LX + 0.22, y + 0.145, 0.5, 0.3, lvl, 10, det_col, True)
    tb(s, LX + 0.72, y + 0.115, 3.35, 0.32, name, 11.5, name_col, True, font=FS)
    tb(s, LX + 4.15, y + 0.155, 3.2, 0.3, detail, 8.5, det_col)
    if badge:
        btxt, bcol = badge
        bw = 0.62 if btxt != "EVOLVED" else 0.82
        rrect(s, LX + LW - bw - 0.16, y + 0.145, bw, 0.27, fill=bcol, radius=0.5)
        tb(s, LX + LW - bw - 0.16, y + 0.195, bw, 0.2, btxt, 7.5, WHITE if btxt != "TODAY" else INK, True, align=PP_ALIGN.CENTER, spc=60)

# consumption loop: agents (L6) feed on ontology (L4) + context graph (L5)
mid = lambda i: ys[i] + LH / 2
bx = LX + LW + 0.24
connector(s, LX + LW + 0.02, mid(1), bx, mid(1), color=TEAL, w=1.75)
connector(s, bx, mid(1), bx, mid(3), color=TEAL, w=1.75)
connector(s, bx, mid(2), LX + LW + 0.06, mid(2), color=TEAL, w=1.75, arrow=True)
connector(s, bx, mid(3), LX + LW + 0.06, mid(3), color=TEAL, w=1.75, arrow=True)

# left rail: value flows up
connector(s, 0.62, ys[6] + LH, 0.62, ys[0] + 0.08, color=GREEN, w=1.75, arrow=True)
vlab = tb(s, -0.66, 4.35, 3.0, 0.25, "VALUE FLOWS UP", 8, RGBColor(0x1E, 0x7A, 0x52), True, align=PP_ALIGN.CENTER, spc=140, wrap=False)
vlab.rotation = 270

# right panel: what changes
tb(s, 10.0, 2.25, 2.85, 0.3, "WHAT CHANGES", 10, FAINT, True, spc=160)
changes = [
    ("Reason over meaning", "Agents work with entities and relationships — not just similar text chunks.", PURPLE),
    ("Lineage by default", "Every claim traces through the graph back to its governed source.", GREEN),
    ("Connect once, reuse everywhere", "New sources snap into the fabric once; every agent benefits instantly.", TEAL),
]
for i, (t, d, col) in enumerate(changes):
    yy = 2.72 + i * 1.32
    dot(s, 10.0, yy + 0.03, 0.16, col)
    tb(s, 10.3, yy - 0.04, 2.6, 0.32, t, 11, INK, True, font=FS)
    tb(s, 10.3, yy + 0.3, 2.6, 0.85, d, 9, GRAY, line_sp=1.15)
footer(s, "Inherits the Enterprise Knowledge Fabric pattern — today's RAG stack already provides layers L1–L3 and L7; L4–L6 are the next investment.", accent=PURPLE)
notes(s, "Takeaway: The future state is not a bigger model — it is smarter knowledge. We add two layers to the fabric we already run: an ontology (shared climate vocabulary) and an enterprise context graph (entities, relationships, lineage). Agents then consume and reason over shared meaning instead of raw text, with traceability built in. Teal loop = agents feeding on L4/L5; today's stack is L1–L3 and L7.")

# ================================================================ SLIDE 15 — TAKEAWAY (dark close)
s = slide_new()
bgr = shp(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, fill=NAVY)
try:
    bgr.fill.gradient()
    stops = bgr.fill.gradient_stops
    stops[0].color.rgb = RGBColor(0x08, 0x22, 0x3C)
    stops[1].color.rgb = RGBColor(0x14, 0x45, 0x63)
    bgr.fill.gradient_angle = 35.0
except Exception:
    pass
cx, cy = 11.15, 5.7
for r in (0.7, 1.25, 1.8):
    ring = shp(s, MSO_SHAPE.OVAL, cx - r, cy - r, 2*r, 2*r, fill=None, line=WHITE, lw=0.9)
    ln = ring.line._get_or_add_ln()
    sf = ln.find(qn('a:solidFill')); clr = sf.find(qn('a:srgbClr'))
    a = clr.makeelement(qn('a:alpha'), {'val': '12000'}); clr.append(a)
dot(s, cx - 0.2, cy - 0.2, 0.4, GREEN)
tb(s, 0.75, 1.7, 6.0, 0.3, "THE TAKEAWAY", 12, RGBColor(0x6F, 0xD0, 0xC2), True, spc=200)
tb(s, 0.72, 2.2, 10.6, 1.9, "Make evidence easier to use —\nwithout losing the ability to verify it.", 34, WHITE, True, font=FS, line_sp=1.05)
tb(s, 0.75, 4.15, 8.6, 1.0,
   "IPCC Climate AI turns a hard-to-navigate report library into a governed, source-grounded research conversation — and the knowledge fabric makes it enterprise-ready.",
   13.5, OFFW, line_sp=1.3)
words = [("DISCOVER", TEAL), ("VERIFY", GREEN), ("ACT", AMBER)]
wx = 0.75
for wtext, wcol in words:
    ww = 0.42 + 0.19 * len(wtext)
    tb(s, wx, 5.6, ww, 0.4, wtext, 16, wcol, True, font=FS, spc=240)
    wx += ww + 0.55
tb(s, 0.75, 6.75, 10.5, 0.3, "IPCC CLIMATE AI  ·  CEO / CIO BRIEFING  ·  v3.0", 9.5, RGBColor(0x7E, 0x95, 0xA8), spc=120)
notes(s, "Takeaway: Close on the promise — speed without losing verifiability. The ask: endorse the pilot gates (quality, security, cost) and the knowledge-fabric investment for the future state.")

# ---------------------------------------------------------------- save
prs.core_properties.title = "IPCC Climate AI — CEO / CIO Briefing v3.0"
prs.core_properties.author = "IPCC AI / GCP Studio"
out = "/home/appadmin/GCP-Studio/IPCC/IPCC_AI_GCP_Overview_v3.0_Final.pptx"
prs.save(out)
print("saved", out)
